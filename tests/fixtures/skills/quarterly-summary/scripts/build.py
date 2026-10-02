"""Build the quarterly summary workbook described in reference/style.md.

Usage: build.py <input.csv> <output.xlsx>

Reads rows of (region, quarter, revenue), writes one "Summary" sheet with a
row per region (Q1-Q4 and Total, largest total first) and an "All regions"
row, and prints the totals as JSON. Standard library plus openpyxl only; it
reads and writes only the two paths it is given.
"""

import csv
import json
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font

QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
MONEY = "#,##0.00"


def main(source: str, target: str) -> None:
    totals = defaultdict(lambda: {quarter: 0.0 for quarter in QUARTERS})
    with open(source, newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        missing = [column for column in ("region", "quarter", "revenue") if column not in (reader.fieldnames or [])]
        if missing:
            print(f"missing column(s): {', '.join(missing)}", file=sys.stderr)
            sys.exit(2)
        for row in reader:
            quarter = row["quarter"].strip().upper()
            if quarter not in QUARTERS:
                print(f"unknown quarter: {row['quarter']!r}", file=sys.stderr)
                sys.exit(2)
            totals[row["region"].strip()][quarter] += float(row["revenue"])

    regions = sorted(totals.items(), key=lambda item: sum(item[1].values()), reverse=True)

    book = Workbook()
    sheet = book.active
    sheet.title = "Summary"
    bold = Font(bold=True)
    sheet.append(["Region", *QUARTERS, "Total"])
    for cell in sheet[1]:
        cell.font = bold
    for region, by_quarter in regions:
        values = [round(by_quarter[quarter], 2) for quarter in QUARTERS]
        sheet.append([region, *values, round(sum(values), 2)])
    column_sums = [round(sum(by_quarter[quarter] for _, by_quarter in regions), 2) for quarter in QUARTERS]
    sheet.append(["All regions", *column_sums, round(sum(column_sums), 2)])
    for cell in sheet[sheet.max_row]:
        cell.font = bold
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=6):
        for cell in row:
            cell.number_format = MONEY
    sheet.freeze_panes = "A2"
    book.save(target)

    print(json.dumps({
        "regions": {region: round(sum(by_quarter.values()), 2) for region, by_quarter in regions},
        "top_region": regions[0][0] if regions else None,
        "grand_total": round(sum(column_sums), 2),
        "sheet": "Summary",
    }))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: build.py <input.csv> <output.xlsx>", file=sys.stderr)
        sys.exit(64)
    main(sys.argv[1], sys.argv[2])
